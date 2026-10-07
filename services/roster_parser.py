"""
직원근무표 엑셀 파서 (Vertex42 월간 달력 템플릿)

주 블록 구조:
  R   : 날짜 행. B열 '사무실근무' 라벨, 요일별 이름 열에 당일 사무실 근무자
  R+1 : '노트북방문' 행. 자유 텍스트 메모
  R+2~: 특이사항. '과정명-담당자,담당자' / '담당자-과정명' / '이름 휴가' / 이름만 있는 연속 행
"""
import re
import logging
from collections import Counter
from datetime import datetime

import openpyxl

from services.excel_parser import (
    INFO_SHEETS, DAY_CONFIG, HOLIDAY_KEYWORDS, NOT_NAME_WORDS, _detect_week_start_rows,
)

logger = logging.getLogger(__name__)

NAME_RE = re.compile(r'^[가-힣]{2,4}$')
HOURS_RE = re.compile(r'^\d+\s*h$', re.IGNORECASE)
PAREN_RE = re.compile(r'\(([^)]*)\)')
TOKEN_SPLIT_RE = re.compile(r'[,，/\s]+')
DELIM_RE = re.compile(r'[-_–]')
LEAVE_WORDS = ('휴가', '연차', '반차', '오전반차', '오후반차', '병가', '경조사')
LEAVE_RE = re.compile(r'^([가-힣]{2,4})\s*(' + '|'.join(LEAVE_WORDS) + r')$')
SKIP_SUBSTRINGS = ('vertex42', 'http://', 'https://')
SKIP_EXACT = {'메모', '사무실근무', '노트북방문'}
# 2~4자 한글이지만 사람 이름이 아닌 단어
NOT_NAME = set(NOT_NAME_WORDS) | set(HOLIDAY_KEYWORDS) | set(LEAVE_WORDS) | {
    '심사', '미팅', '회의', '수업', '교육', '방문', '예정', '출장', '재택',
    '타운홀', '벡스코', '숙대', '마산대', '세종시', '공휴일', '반납', '멘토링',
}

OFFICE_LABEL = '사무실근무'
LAPTOP_LABEL = '노트북방문'


def _clean(text):
    return re.sub(r'\s+', ' ', str(text).replace('\n', ' ')).strip()


def _is_noise(text):
    low = text.lower()
    if not text or text in SKIP_EXACT or HOURS_RE.match(text):
        return True
    return any(s in low for s in SKIP_SUBSTRINGS)


def _is_name(token):
    return bool(NAME_RE.match(token)) and token not in NOT_NAME


def _split_names(text):
    """토큰이 전부 한국어 이름이면 이름 리스트, 아니면 None"""
    tokens = [t for t in TOKEN_SPLIT_RE.split(text.strip()) if t]
    if tokens and all(_is_name(t) for t in tokens):
        return tokens
    return None


def _extract_note(text):
    """괄호 메모 분리 → (괄호 제거 텍스트, 메모)"""
    notes = [m.strip() for m in PAREN_RE.findall(text) if m.strip()]
    stripped = _clean(PAREN_RE.sub(' ', text))
    return stripped, ' / '.join(notes)


def _parse_event_text(text, known):
    """'과정명-담당자' 또는 '담당자-과정명' → (title, people, note)"""
    body, note = _extract_note(text)
    for m in DELIM_RE.finditer(body):
        left, right = body[:m.start()].strip(), body[m.end():].strip()
        if not left or not right:
            continue
        ln, rn = _split_names(left), _split_names(right)
        if rn and not ln:
            return left, rn, note
        if ln and not rn:
            return right, ln, note
        if ln and rn:
            # 양쪽 다 이름 같으면 알려진 직원이 더 많은 쪽을 담당자로 (동률이면 '과정명-담당자' 관례)
            lk, rk = sum(n in known for n in ln), sum(n in known for n in rn)
            return (right, ln, note) if lk > rk else (left, rn, note)
    return body, [], note


def _sheet_month(ws):
    """시트가 나타내는 (year, month). A1 우선, 없으면 날짜 셀 최빈값"""
    a1 = ws.cell(row=1, column=1).value
    if isinstance(a1, datetime) and a1.year >= 2020:
        return a1.year, a1.month
    counter = Counter()
    for row in range(8, min(ws.max_row, 55) + 1):
        for date_col, _ in DAY_CONFIG:
            v = ws.cell(row=row, column=date_col).value
            if isinstance(v, datetime) and v.year >= 2020:
                counter[(v.year, v.month)] += 1
    return counter.most_common(1)[0][0] if counter else None


def get_roster_sheet_info(filepath):
    """시트별 (이름, 연, 월) 목록 — 업로드 UI 안내용"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    info = []
    for ws in wb.worksheets:
        if ws.title in INFO_SHEETS:
            continue
        ym = _sheet_month(ws)
        info.append({
            "name": ws.title,
            "year": ym[0] if ym else None,
            "month": ym[1] if ym else None,
        })
    wb.close()
    return info


def _cell_texts(ws, row, cols):
    """지정 행/열들의 문자열 값 (중복 제거, 순서 유지)"""
    seen = []
    for col in cols:
        v = ws.cell(row=row, column=col).value
        if v is None or isinstance(v, (datetime, int, float)):
            continue
        text = str(v).strip()
        if text and text not in seen:
            seen.append(text)
    return seen


def _week_blocks(ws):
    """(시작행, 끝행) 목록"""
    week_rows = _detect_week_start_rows(ws)
    ends = [r - 1 for r in week_rows[1:]] + [min(ws.max_row, week_rows[-1] + 10)]
    return list(zip(week_rows, ends))


def _block_dates(ws, start_row, year, month):
    """블록 날짜 행에서 (date, date_col, name_col) — 해당 월만"""
    out = []
    for date_col, name_col in DAY_CONFIG:
        v = ws.cell(row=start_row, column=date_col).value
        if isinstance(v, datetime) and (v.year, v.month) == (year, month):
            out.append((v, date_col, name_col))
    return out


def _collect_known_names(sheets):
    """1차 패스: 사무실근무 행의 단일 이름들을 직원 명단으로 수집"""
    known = set()
    for ws, year, month in sheets:
        for start_row, _ in _week_blocks(ws):
            for _, _, name_col in _block_dates(ws, start_row, year, month):
                for text in _cell_texts(ws, start_row, [name_col]):
                    names = _split_names(_clean(text))
                    if names:
                        known.update(names)
    return known


def _make_entry(date, category, title, people, note, raw, sheet):
    return {
        "date": date.strftime('%Y-%m-%d'),
        "category": category,
        "title": title,
        "people": people,
        "note": note,
        "raw_text": raw,
        "source_sheet": sheet,
    }


def _classify(text, known):
    """특이사항 텍스트 → (category, title, people, note)"""
    cleaned = _clean(text)
    m = LEAVE_RE.match(cleaned)
    if m:
        return 'leave', m.group(2), [m.group(1)], ''
    title, people, note = _parse_event_text(cleaned, known)
    if title in LEAVE_WORDS and people:
        return 'leave', title, people, note
    if not people and any(kw in cleaned for kw in HOLIDAY_KEYWORDS):
        return 'memo', title, [], note
    return 'event', title, people, note


def _parse_block(ws, start_row, end_row, year, month, sheet_name, known):
    entries = []
    for date_cell, date_col, name_col in _block_dates(ws, start_row, year, month):
        cols = (date_col, name_col)

        # R: 사무실 근무자 ('이름', '이름-업무', '대체휴일' 등)
        for text in _cell_texts(ws, start_row, [name_col]):
            cleaned = _clean(text)
            if _is_noise(cleaned):
                continue
            names = _split_names(cleaned)
            if names:
                entries.append(_make_entry(date_cell, 'office', OFFICE_LABEL, names, '', text, sheet_name))
                continue
            body, note = _extract_note(cleaned)
            names = _split_names(body) if body else None
            if names:
                entries.append(_make_entry(date_cell, 'office', OFFICE_LABEL, names, note, text, sheet_name))
                continue
            category, title, people, note = _classify(text, known)
            if category == 'event' and people:
                entries.append(_make_entry(date_cell, 'office', OFFICE_LABEL, people,
                                           ' / '.join(x for x in (title, note) if x), text, sheet_name))
            elif category == 'event':
                entries.append(_make_entry(date_cell, 'office', OFFICE_LABEL, [], cleaned, text, sheet_name))
            else:
                entries.append(_make_entry(date_cell, category, title, people, note, text, sheet_name))

        # R+1: 노트북 방문 메모 (직원 일정이 적혀 있으면 event로)
        for text in _cell_texts(ws, start_row + 1, cols):
            cleaned = _clean(text)
            if _is_noise(cleaned):
                continue
            category, title, people, note = _classify(text, known)
            if category != 'event' or any(p in known for p in people):
                entries.append(_make_entry(date_cell, category, title, people, note, text, sheet_name))
            else:
                entries.append(_make_entry(date_cell, 'laptop', LAPTOP_LABEL, [], cleaned, text, sheet_name))

        # R+2~: 특이사항
        last_event = None
        for row in range(start_row + 2, end_row + 1):
            for text in _cell_texts(ws, row, cols):
                cleaned = _clean(text)
                if _is_noise(cleaned):
                    continue

                names_only = _split_names(cleaned)
                # 단일 토큰은 이름이 아닐 수 있어 직전 일정이 있거나 알려진 직원일 때만 담당자로 취급
                if names_only and (last_event is not None or len(names_only) >= 2
                                   or names_only[0] in known):
                    if last_event is not None:
                        for n in names_only:
                            if n not in last_event['people']:
                                last_event['people'].append(n)
                        last_event['raw_text'] += '\n' + text
                    else:
                        last_event = _make_entry(date_cell, 'event', '', names_only, '', text, sheet_name)
                        entries.append(last_event)
                    continue

                category, title, people, note = _classify(text, known)
                entry = _make_entry(date_cell, category, title, people, note, text, sheet_name)
                entries.append(entry)
                last_event = entry if category == 'event' else None

    return entries


def parse_roster(filepath, selected_sheets):
    """직원근무표 엑셀 → (RosterEntry dict 리스트, [(year, month), ...])"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheets = []
    for sheet_name in selected_sheets:
        if sheet_name in INFO_SHEETS or sheet_name not in wb.sheetnames:
            logger.warning(f"시트 '{sheet_name}' 건너뜀")
            continue
        ws = wb[sheet_name]
        ym = _sheet_month(ws)
        if not ym:
            logger.warning(f"시트 '{sheet_name}': 월을 판별할 수 없음")
            continue
        sheets.append((ws, ym[0], ym[1]))

    known = _collect_known_names(sheets)
    logger.info(f"직원 명단(사무실근무 기준): {sorted(known)}")

    entries = []
    for ws, year, month in sheets:
        before = len(entries)
        for start_row, end_row in _week_blocks(ws):
            entries.extend(_parse_block(ws, start_row, end_row, year, month, ws.title, known))
        logger.info(f"시트 '{ws.title}' ({year}-{month:02d}): {len(entries) - before}건")

    wb.close()
    months = sorted({(y, m) for _, y, m in sheets})
    return entries, months
