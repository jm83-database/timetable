"""수업 장소('?' 파싱, 입력 검증, 메일 문구) 테스트"""
from datetime import datetime

import openpyxl
import pytest

from routes import _resolve_location
from services.excel_parser import _location_fields, _split_location, parse_timetable
from services.reminder_sender import message_for


def _write_timetable(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '1기'
    ws.cell(row=10, column=3, value=datetime(2026, 11, 4))
    ws.cell(row=10, column=4, value='AI기본 ?서울특별시교육청 OO로 1')
    ws.cell(row=11, column=4, value='강명호')
    ws.cell(row=10, column=5, value=datetime(2026, 11, 5))
    ws.cell(row=10, column=6, value='AI심화')
    ws.cell(row=11, column=6, value='인선미')
    ws.cell(row=12, column=6, value='?온라인')
    wb.save(path)


def test_split_location_handles_halfwidth_and_fullwidth_marks():
    assert _split_location('AI기본 ?서울 OO로') == ('AI기본', '서울 OO로')
    assert _split_location('AI기본 ？온라인') == ('AI기본', '온라인')
    assert _split_location('강명호') == ('강명호', None)


def test_location_fields_defaults_to_online():
    assert _location_fields(None) == {'location_type': 'online', 'location': ''}
    assert _location_fields('온라인') == {'location_type': 'online', 'location': ''}
    assert _location_fields('서울특별시교육청') == {'location_type': 'offline', 'location': '서울특별시교육청'}


def test_parse_timetable_extracts_location_per_entry(tmp_path):
    path = tmp_path / 'timetable.xlsx'
    _write_timetable(path)

    entries = {e['class_name']: e for e in parse_timetable(str(path), ['1기'])}

    offline = entries['AI기본']
    assert offline['instructor'] == '강명호'
    assert offline['location_type'] == 'offline'
    assert offline['location'] == '서울특별시교육청 OO로 1'

    online = entries['AI심화']
    assert online['instructor'] == '인선미'
    assert online['location_type'] == 'online'
    assert online['location'] == ''


def test_resolve_location_validates_offline_input():
    assert _resolve_location({'location_type': 'online', 'location': '무시'}) == {
        'location_type': 'online', 'location': ''}
    assert _resolve_location({'location_type': 'offline', 'location': ' 서울교육청 '}) == {
        'location_type': 'offline', 'location': '서울교육청'}
    with pytest.raises(ValueError):
        _resolve_location({'location_type': 'offline', 'location': '   '})
    with pytest.raises(ValueError):
        _resolve_location({'location_type': 'offline', 'location': 'a' * 201})
    with pytest.raises(ValueError):
        _resolve_location({'location_type': 'hybrid'})


def test_mail_body_includes_online_or_offline_location():
    course = {'name': '교육'}
    base = {'class_name': '수업', 'date': '2026-11-04', 'start_time': '09:00',
            'end_time': '18:00', 'instructor': '강사'}

    _, online_body = message_for(course, {**base, 'location_type': 'online', 'location': ''}, 7)
    assert '장소: 온라인' in online_body

    _, offline_body = message_for(course, {**base, 'location_type': 'offline',
                                           'location': '서울특별시교육청'}, 3)
    assert '장소: 서울특별시교육청' in offline_body

    _, legacy_body = message_for(course, base, 7)
    assert '장소: 온라인' in legacy_body
